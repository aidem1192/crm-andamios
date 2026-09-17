"""
Importa clientes desde CLIENTES_LISTOS_PARA_IMPORTAR.xlsx al CRM.
Uso: python importar_clientes.py

El script se conecta directamente a la base de datos (sin pasar por la API)
para mayor velocidad al importar 2,000+ registros.
"""
import sys
import os
import re

# Ruta al Excel generado
XLSX = r"G:\Mi unidad\ANDAMIOS Y DERIVADOS\ANDAMIOS\DOCUMENTACION INTERNA\BASE DE DATOS\CLIENTES_LISTOS_PARA_IMPORTAR.xlsx"

# Configurar path para importar los modelos del CRM
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault("DATABASE_URL", "sqlite:///./data/andamios.db")

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl no instalado. Ejecuta: pip install openpyxl")
    sys.exit(1)

from app.database import SessionLocal, Cliente, ReferenciaCliente, init_db

init_db()

print(f"\n{'='*60}")
print("  IMPORTADOR DE CLIENTES — CRM Andamios y Derivados del Norte")
print(f"{'='*60}")
print(f"Fuente: {XLSX}\n")

if not os.path.exists(XLSX):
    print(f"ERROR: No se encontro el archivo:\n  {XLSX}")
    sys.exit(1)

wb = openpyxl.load_workbook(XLSX)
ws = wb["Clientes"]

db = SessionLocal()

creados = 0
omitidos = 0
lista_negra_marcados = 0
errores = []

total_filas = ws.max_row - 1
print(f"Total de clientes en archivo: {total_filas}")

# Obtener nombres ya existentes para no duplicar
existentes = {c.nombre.upper().strip() for c in db.query(Cliente).all()}
print(f"Clientes ya en la BD: {len(existentes)}")
print(f"\nImportando", end="", flush=True)

for ri, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 1):
    nombre  = str(row[0] or "").strip()
    tel     = str(row[1] or "").strip()
    dom     = str(row[2] or "").strip()
    r1_nom  = str(row[3] or "").strip()
    r1_tel  = str(row[4] or "").strip()
    r2_nom  = str(row[5] or "").strip()
    r2_tel  = str(row[6] or "").strip()
    es_ln   = str(row[7] or "").strip().upper() == "SI"
    motivo  = str(row[8] or "").strip()

    if not nombre or nombre == "None":
        continue

    # Saltar si ya existe
    if nombre.upper().strip() in existentes:
        omitidos += 1
        continue

    # Normalizar teléfono — guardar solo dígitos si son 10
    tel_clean = re.sub(r"[^\d]", "", tel)
    tel_final = tel_clean if len(tel_clean) == 10 else tel[:30] if tel else None

    try:
        cliente = Cliente(
            nombre=nombre,
            telefono=tel_final,
            domicilio=dom or None,
            lista_negra=es_ln,
            motivo_lista_negra=motivo if es_ln else None,
        )
        db.add(cliente)
        db.flush()  # obtener el id

        # Referencias
        if r1_nom and r1_nom.lower() not in ("none", ""):
            r1_tel_d = re.sub(r"[^\d]", "", r1_tel)[:10]
            db.add(ReferenciaCliente(
                cliente_id=cliente.id,
                nombre=r1_nom[:200],
                telefono=r1_tel_d or None,
            ))

        if r2_nom and r2_nom.lower() not in ("none", ""):
            r2_tel_d = re.sub(r"[^\d]", "", r2_tel)[:10]
            db.add(ReferenciaCliente(
                cliente_id=cliente.id,
                nombre=r2_nom[:200],
                telefono=r2_tel_d or None,
            ))

        existentes.add(nombre.upper().strip())
        creados += 1
        if es_ln:
            lista_negra_marcados += 1

        # Commit cada 100 registros
        if creados % 100 == 0:
            db.commit()
            print(".", end="", flush=True)

    except Exception as e:
        db.rollback()
        errores.append(f"Fila {ri+1}: {nombre} → {e}")

# Commit final
try:
    db.commit()
except Exception as e:
    db.rollback()
    errores.append(f"Commit final: {e}")

db.close()

print(f"\n\n{'='*60}")
print(f"  RESULTADO")
print(f"{'='*60}")
print(f"  ✓ Clientes creados:          {creados}")
print(f"  ⛔ En lista negra marcados:  {lista_negra_marcados}")
print(f"  ↷ Ya existían (omitidos):   {omitidos}")
if errores:
    print(f"  ✗ Errores:                  {len(errores)}")
    for e in errores[:10]:
        print(f"      {e}")
    if len(errores) > 10:
        print(f"      ... y {len(errores)-10} más")
print(f"{'='*60}\n")
