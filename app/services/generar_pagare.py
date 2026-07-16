"""
Genera un Pagaré conforme a la Ley General de Títulos y Operaciones de Crédito
(LGTOC), Artículos 170-174, para garantizar la devolución del material rentado.

Elementos obligatorios (Art. 170 LGTOC):
  I.   Mención de ser PAGARÉ en el texto del documento.
  II.  Promesa incondicional de pagar una suma determinada.
  III. Nombre de la persona a quien se pagará (beneficiaria).
  IV.  Época y lugar del pago.
  V.   Fecha y lugar de suscripción.
  VI.  Firma del suscriptor.
"""

from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from app.services.numero_a_letra import numero_a_letra

OUTPUT_DIR = Path(__file__).parent.parent.parent / "data" / "contratos"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

MESES_ES = [
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
]

EMPRESA = "ANDAMIOS Y DERIVADOS DEL NORTE, S.A. DE C.V."
EMPRESA_CORTA = "ANDAMIOS Y DERIVADOS DEL NORTE"
DOMICILIO_EMPRESA = "Chihuahua, Chihuahua, México"


def _fecha_larga(d: date) -> str:
    return f"{d.day} de {MESES_ES[d.month - 1]} de {d.year}"


def _set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def _set_cell_borders(cell, top=True, bottom=True, left=True, right=True):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    sides = {
        'top': top, 'bottom': bottom, 'left': left, 'right': right
    }
    for side, active in sides.items():
        el = OxmlElement(f'w:{side}')
        if active:
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), '6')
            el.set(qn('w:space'), '0')
            el.set(qn('w:color'), '1a1a1a')
        else:
            el.set(qn('w:val'), 'none')
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _para(doc, text='', bold=False, italic=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT,
           color=None, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)
    return p


def _add_run(para, text, bold=False, italic=False, size=11, color=None, underline=False):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.underline = underline
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def generar_pagare(contrato_data: dict) -> Path:
    """
    contrato_data keys:
      folio, cliente_nombre, cliente_domicilio, cliente_rfc, cliente_telefono,
      lineas: [{material, cantidad, valor_convencional}]
      fecha_inicio (date), lugar_obra, lugar_celebracion
    """
    folio = contrato_data["folio"]
    output_path = OUTPUT_DIR / f"pagare_{folio}.docx"

    cliente     = contrato_data["cliente_nombre"].upper()
    domicilio   = contrato_data.get("cliente_domicilio", "").upper()
    rfc         = contrato_data.get("cliente_rfc", "").upper()
    telefono    = contrato_data.get("cliente_telefono", "")
    lineas      = contrato_data.get("lineas", [])
    fecha_ini   = contrato_data["fecha_inicio"]
    lugar_cel   = contrato_data.get("lugar_celebracion", "CHIHUAHUA, CHIHUAHUA")

    # Calcular valor convencional total (reposición del material)
    total_conv = sum(
        float(l.get("valor_convencional", 0)) * int(l.get("cantidad", 1))
        for l in lineas
    )
    letra = numero_a_letra(total_conv)

    doc = Document()

    # ── Márgenes ──
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)
        section.page_width    = Cm(21.59)   # Carta
        section.page_height   = Cm(27.94)

    # ── Encabezado empresa ──
    header_table = doc.add_table(rows=1, cols=2)
    header_table.style = 'Table Grid'
    _set_cell_bg(header_table.rows[0].cells[0], 'FF6B00')
    _set_cell_bg(header_table.rows[0].cells[1], 'FFFFFF')
    header_table.rows[0].cells[0].width = Cm(10)
    header_table.rows[0].cells[1].width = Cm(8)

    c0 = header_table.rows[0].cells[0]
    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('ANDAMIOS Y DERIVADOS')
    r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor(255, 255, 255)
    p2 = c0.add_paragraph('DEL NORTE')
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.runs[0] if p2.runs else p2.add_run('DEL NORTE')
    r2.bold = True; r2.font.size = Pt(14); r2.font.color.rgb = RGBColor(255, 255, 255)

    c1 = header_table.rows[0].cells[1]
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = c1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_run(p, f'Folio: {folio}', bold=True, size=10)
    p2 = c1.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_run(p2, f'Fecha: {_fecha_larga(fecha_ini)}', size=9, italic=True)

    doc.add_paragraph()

    # ── Título ──
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_before = Pt(4)
    titulo.paragraph_format.space_after = Pt(2)
    r = titulo.add_run('P A G A R É')
    r.bold = True; r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(255, 107, 0)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(10)
    _add_run(sub, f'(Garantía por reposición de material rentado — Contrato {folio})',
             italic=True, size=9, color=(120, 120, 120))

    # ── Monto en caja destacada ──
    monto_table = doc.add_table(rows=1, cols=3)
    monto_table.style = 'Table Grid'
    monto_table.rows[0].height = Cm(1.4)

    _set_cell_bg(monto_table.rows[0].cells[0], 'F5F5F5')
    _set_cell_bg(monto_table.rows[0].cells[1], 'FF6B00')
    _set_cell_bg(monto_table.rows[0].cells[2], 'F5F5F5')

    mc0 = monto_table.rows[0].cells[0]
    mc0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = mc0.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(p, 'IMPORTE', bold=True, size=9, color=(100, 100, 100))

    mc1 = monto_table.rows[0].cells[1]
    mc1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = mc1.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(p, f'$ {total_conv:,.2f} M.N.', bold=True, size=16, color=(255, 255, 255))

    mc2 = monto_table.rows[0].cells[2]
    mc2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = mc2.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(p, 'PESOS M.N.', bold=True, size=9, color=(100, 100, 100))

    doc.add_paragraph()

    # ── Cuerpo del pagaré ──
    cuerpo = doc.add_paragraph()
    cuerpo.paragraph_format.space_after = Pt(10)
    cuerpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    _add_run(cuerpo, f'En {lugar_cel}, a {_fecha_larga(fecha_ini)}. — Yo, ')
    _add_run(cuerpo, cliente, bold=True, underline=True)
    _add_run(cuerpo, ', con domicilio en ')
    _add_run(cuerpo, domicilio or '____________________________', bold=True)
    if rfc:
        _add_run(cuerpo, f', RFC: {rfc}')
    _add_run(cuerpo,
        ', en mi carácter de cliente y suscriptor del presente título de crédito, '
        'me comprometo ')
    _add_run(cuerpo, 'incondicionalmente', bold=True, underline=True)
    _add_run(cuerpo, ' a pagar a la orden de ')
    _add_run(cuerpo, EMPRESA, bold=True)
    _add_run(cuerpo, f', o a quien sus derechos represente, la cantidad de ')
    _add_run(cuerpo, f'$ {total_conv:,.2f} ({letra})', bold=True)
    _add_run(cuerpo, ' en moneda nacional.')

    cuerpo2 = doc.add_paragraph()
    cuerpo2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cuerpo2.paragraph_format.space_after = Pt(10)
    _add_run(cuerpo2,
        'El presente Pagaré ampara el valor convencional de reposición del material '
        'entregado en renta conforme al Contrato ')
    _add_run(cuerpo2, folio, bold=True)
    _add_run(cuerpo2, '. Este título se tornará exigible de manera inmediata en caso de '
        'que el material no sea devuelto en su totalidad y en las mismas condiciones '
        'en que fue entregado, o ante cualquier incumplimiento del contrato de renta.')

    cuerpo3 = doc.add_paragraph()
    cuerpo3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cuerpo3.paragraph_format.space_after = Pt(10)
    _add_run(cuerpo3,
        'El pago deberá realizarse en el domicilio de ')
    _add_run(cuerpo3, EMPRESA_CORTA, bold=True)
    _add_run(cuerpo3, f', ubicado en {DOMICILIO_EMPRESA}, o en el lugar que '
        'designe el tenedor legítimo. En caso de mora, el suscriptor pagará adicionalmente '
        'los intereses moratorios a la tasa del ')
    _add_run(cuerpo3, '6% mensual', bold=True)
    _add_run(cuerpo3,
        ', así como los gastos y costas de cobranza judicial o extrajudicial. '
        'El presente pagaré se rige por la ')
    _add_run(cuerpo3,
        'Ley General de Títulos y Operaciones de Crédito (LGTOC), Arts. 170-174.',
        italic=True)

    # ── Tabla de materiales ──
    _para(doc, 'Descripción del material garantizado:', bold=True, size=10, space_before=8, space_after=4)

    mat_table = doc.add_table(rows=1, cols=4)
    mat_table.style = 'Table Grid'

    hdrs = ['MATERIAL', 'CANTIDAD', 'VALOR UNIT.', 'VALOR TOTAL']
    widths = [Cm(7.5), Cm(2.5), Cm(3), Cm(3)]
    for i, (h, w) in enumerate(zip(hdrs, widths)):
        cell = mat_table.rows[0].cells[i]
        cell.width = w
        _set_cell_bg(cell, '2D2D2D')
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_run(p, h, bold=True, size=9, color=(255, 255, 255))

    for l in lineas:
        nombre   = str(l.get('material', '')).upper()
        cantidad = int(l.get('cantidad', 1))
        val_unit = float(l.get('valor_convencional', 0))
        val_tot  = val_unit * cantidad

        row = mat_table.add_row()
        row.cells[0].width = widths[0]; row.cells[1].width = widths[1]
        row.cells[2].width = widths[2]; row.cells[3].width = widths[3]

        def _celd(idx, txt, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False):
            p = row.cells[idx].paragraphs[0]; p.alignment = align
            _add_run(p, txt, size=9, bold=bold)

        _celd(0, nombre)
        _celd(1, str(cantidad), WD_ALIGN_PARAGRAPH.CENTER)
        _celd(2, f'$ {val_unit:,.2f}', WD_ALIGN_PARAGRAPH.RIGHT)
        _celd(3, f'$ {val_tot:,.2f}', WD_ALIGN_PARAGRAPH.RIGHT, bold=True)

    # Fila de total
    tot_row = mat_table.add_row()
    _set_cell_bg(tot_row.cells[0], 'F5F5F5')
    _set_cell_bg(tot_row.cells[2], 'F5F5F5')
    _set_cell_bg(tot_row.cells[3], 'FF6B00')

    p = tot_row.cells[2].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_run(p, 'TOTAL:', bold=True, size=9)
    p = tot_row.cells[3].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_run(p, f'$ {total_conv:,.2f}', bold=True, size=10, color=(255, 255, 255))

    # ── Mención lugar de obra ──
    if contrato_data.get("lugar_obra"):
        _para(doc, f'Lugar de entrega del material: {contrato_data["lugar_obra"]}',
              italic=True, size=9, space_before=8)

    doc.add_paragraph()

    # ── Firma ──
    firma_table = doc.add_table(rows=3, cols=2)
    firma_table.style = 'Table Grid'

    # Celda izquierda: datos del suscriptor
    _set_cell_bg(firma_table.rows[0].cells[0], 'F9F9F9')
    p = firma_table.rows[0].cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _add_run(p, 'SUSCRIPTOR (CLIENTE):', bold=True, size=9, color=(80, 80, 80))

    firma_table.rows[1].cells[0].merge(firma_table.rows[1].cells[0])
    p = firma_table.rows[1].cells[0].paragraphs[0]
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(p, '________________________________________', size=10)

    p = firma_table.rows[2].cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(p, cliente, bold=True, size=9)
    p2 = firma_table.rows[2].cells[0].add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(p2, f'Tel: {telefono}  |  RFC: {rfc}' if rfc else f'Tel: {telefono}', size=8, italic=True)

    # Celda derecha: testigo / huella
    _set_cell_bg(firma_table.rows[0].cells[1], 'F9F9F9')
    p = firma_table.rows[0].cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _add_run(p, 'HUELLA DACTILAR / TESTIGO:', bold=True, size=9, color=(80, 80, 80))

    p = firma_table.rows[1].cells[1].paragraphs[0]
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(p, '________________________________________', size=10)

    p = firma_table.rows[2].cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(p, 'Nombre y firma del testigo', size=9, italic=True)

    doc.add_paragraph()

    # ── Pie legal ──
    pie = doc.add_paragraph()
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_run(pie,
        'Este documento es un título de crédito conforme a los Arts. 170-174 de la LGTOC. '
        'Su sola firma constituye reconocimiento de deuda líquida y exigible. '
        f'Emitido por {EMPRESA}.',
        size=7, italic=True, color=(150, 150, 150))

    doc.save(str(output_path))
    return output_path
