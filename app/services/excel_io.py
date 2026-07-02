"""Generación de plantillas Excel y lectura de archivos importados."""
import io
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="F97316", end_color="F97316", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _build_workbook(headers: list[str], ejemplo: list, notas: list[str], anchos: list[int] = None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    for col, titulo in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=titulo)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")

    if ejemplo:
        for col, val in enumerate(ejemplo, start=1):
            ws.cell(row=2, column=col, value=val)

    for col, ancho in enumerate(anchos or [20] * len(headers), start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho

    ws.freeze_panes = "A2"

    if notas:
        ws2 = wb.create_sheet("Instrucciones")
        ws2.column_dimensions["A"].width = 90
        for i, nota in enumerate(notas, start=1):
            ws2.cell(row=i, column=1, value=nota)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def plantilla_clientes() -> io.BytesIO:
    headers = ["nombre", "telefono", "domicilio", "rfc", "email",
               "ref1_nombre", "ref1_telefono", "ref1_direccion",
               "ref2_nombre", "ref2_telefono", "ref2_direccion"]
    ejemplo = ["Gilberto Loya Castillo", "614-192-96-47", "C. 62 #2000, Col. Cerro de la Cruz", "LOCG800101ABC",
               "gilberto@correo.com", "Joel González", "614-486-25-04", "C. Mina de San Carlos #17548",
               "", "", ""]
    notas = [
        "Instrucciones para importar clientes",
        "",
        "1. No modifiques ni borres los encabezados de la fila 1 (hoja 'Datos').",
        "2. La fila 2 es un ejemplo; puedes borrarla o sobreescribirla.",
        "3. 'nombre' es el único campo obligatorio.",
        "4. Las referencias (ref1_*, ref2_*) son opcionales; se admiten hasta 2 por cliente.",
        "5. Guarda el archivo en formato .xlsx y súbelo en la sección Clientes.",
    ]
    anchos = [26, 16, 32, 16, 24, 20, 16, 28, 20, 16, 28]
    return _build_workbook(headers, ejemplo, notas, anchos)


def plantilla_materiales() -> io.BytesIO:
    headers = ["nombre", "tipo", "numero_serie", "estado", "ubicacion",
               "precio_renta_dia", "precio_venta", "descripcion", "stock"]
    ejemplo = ["Marco de andamio 90x190", "pieza", "", "bueno", "bodega", 8.0, 1200.0,
               "Marco galvanizado", 20]
    notas = [
        "Instrucciones para importar materiales",
        "",
        "1. No modifiques ni borres los encabezados de la fila 1 (hoja 'Datos').",
        "2. La fila 2 es un ejemplo; puedes borrarla o sobreescribirla.",
        "3. Campos obligatorios: 'nombre' y 'precio_renta_dia'.",
        "4. 'tipo' debe ser: pieza, kit_fijo o kit_personalizable (usa 'pieza' si no estás seguro).",
        "5. 'estado' debe ser: bueno, danado o en_reparacion.",
        "6. 'ubicacion' debe ser: bodega, en_obra o en_transito.",
        "7. 'stock' solo aplica a materiales tipo 'pieza' (déjalo en 0 para kits).",
        "8. Los kits necesitan configurar sus componentes manualmente desde la app despues de importar.",
        "9. Guarda el archivo en formato .xlsx y súbelo en la sección Materiales.",
    ]
    anchos = [28, 20, 16, 16, 14, 16, 14, 30, 10]
    return _build_workbook(headers, ejemplo, notas, anchos)


def leer_filas(archivo_bytes: bytes) -> list[dict]:
    """Lee la hoja 'Datos' (o la primera hoja) y regresa una lista de dicts por fila,
    usando los encabezados de la fila 1 tal como aparecen en el archivo."""
    wb = load_workbook(io.BytesIO(archivo_bytes), data_only=True)
    ws = wb["Datos"] if "Datos" in wb.sheetnames else wb.worksheets[0]

    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []

    encabezados = [str(h).strip() if h is not None else "" for h in filas[0]]
    registros = []
    for fila in filas[1:]:
        if fila is None or all(v is None or str(v).strip() == "" for v in fila):
            continue
        registro = {}
        for i, h in enumerate(encabezados):
            if not h:
                continue
            registro[h] = fila[i] if i < len(fila) else None
        registros.append(registro)
    return registros


def limpiar_texto(v) -> str | None:
    """Normaliza un valor de celda a string limpio, o None si está vacío."""
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None
